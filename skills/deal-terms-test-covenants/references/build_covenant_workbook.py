"""
Canonical Excel builder for the test-covenants skill.

Produces the standard Covenant Test workbook. The skill parses
its render into the `data` dict below and calls build(data, out_path). Sheets:
  1. Covenant Package  — by-tranche package table + cov-lite warning + summary
  2. Spread            — one spread block per binding covenant, Pass/Fail color-coded
  3. Stress            — one breakeven + stress grid per binding covenant, cells color-coded

INPUT SHAPE:

data = {
  "title": "Covenant Test — [Deal Name] ([Company])",
  "meta":  "USD in millions · Data confidence: … · Generated 2026-08-31",
  "stale_note": "…" | None,            # rendered as a caution line if present
  "ebitda_basis": "Adjusted EBITDA as the model determines it. …",
  "covlite_warning": "…" | None,       # amber warning banner
  "package": [                          # one row per facility
     {"facility":"Revolving Facility","regime":"Springing","metric":"First Lien Net Leverage",
      "level":"9.77x [Assumed]","frequency":"Quarterly, if sprung",
      "trigger":"Utilization > 40% of commitments","cure":"Yes — max 5 post-close"},
     {"facility":"Term Loan","regime":"Cov-lite","metric":"None","level":"—",
      "frequency":"—","trigger":"—","cure":"—"},
  ],
  "package_note": "No covenant level disclosed; 9.77x assumed …" | None,
  "summary": {"compliance":"…","key_risk":"…","springing":"…"},
  "spreads": [                          # one per BINDING covenant
     {"facility":"Revolving Facility","metric":"First Lien Net Leverage",
      "note":"Springing covenant: modeled revolver undrawn, so dormant; shown as-if-sprung." | None,
      "rows":[("LTM Oct-24","7.24x","9.77x","Pass","2.53x"), …]},  # (period, metric, threshold, status, headroom)
  ],                                    # status ∈ {Pass, Tight, Breach}
  "covlite_note": "Term Loan cov-lite …" | None,
  "stress": [                           # one per BINDING covenant
     {"facility":"Revolving Facility","metric":"First Lien Net Leverage","threshold":9.77,
      "rows":[("2025E","$63.1","$319.3",5.06,5.62,6.32,7.23,"-48.2%"), …]},
       # (year, ebitda_str, debt_str, base, m10, m20, m30, breakeven_str) — leverage as floats
  ],
  "web_url": "https://…",
}
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

NAVY="1F2933"; SLATE="3D4A5C"; GREY="6B7280"
GREEN_F="D6EFDF"; GREEN_T="1E6B3A"; AMBER_F="FCF3CF"; AMBER_T="7D6608"; RED_F="F5C6CB"; RED_T="922B21"
_thin=Side(style="thin", color="D9DDE3"); BORDER=Border(_thin,_thin,_thin,_thin)

def _f(sz=10,b=False,color=NAVY,italic=False): return Font(name="Arial",size=sz,bold=b,color=color,italic=italic)
def _fill(c): return PatternFill("solid",fgColor=c)
WRAP=Alignment(wrap_text=True,vertical="top"); CENTER=Alignment(horizontal="center",vertical="center",wrap_text=True)
LEFT=Alignment(horizontal="left",vertical="center"); LEFTW=Alignment(horizontal="left",vertical="center",wrap_text=True)

def _status_colors(status):
    st=(status or "").lower()
    if "breach" in st or "fail" in st or "🔴" in status: return RED_F,RED_T
    if "tight" in st or "⚠" in status:                    return AMBER_F,AMBER_T
    return GREEN_F,GREEN_T

def _lev_colors(value, threshold):
    if threshold is None or value is None: return None,None
    if value > threshold:                  return RED_F,RED_T       # breach
    if threshold - value < 0.25:           return AMBER_F,AMBER_T   # tight
    return GREEN_F,GREEN_T                                          # pass

def _banner(ws,r,text,span):
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=span)
    c=ws.cell(r,1,text); c.font=_f(11,True,"FFFFFF"); c.fill=_fill(SLATE)
    c.alignment=Alignment(horizontal="left",vertical="center",indent=1); ws.row_dimensions[r].height=22

def _title(ws,r,text,sub,span):
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=span)
    ws.cell(r,1,text).font=_f(14,True,NAVY); ws.row_dimensions[r].height=26
    ws.merge_cells(start_row=r+1,start_column=1,end_row=r+1,end_column=span)
    ws.cell(r+1,1,sub).font=_f(9,False,GREY,italic=True)

def _para(ws,r,text,span,ht=40,font=None,fillc=None):
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=span)
    c=ws.cell(r,1,text); c.font=font or _f(10); c.alignment=WRAP
    if fillc: c.fill=_fill(fillc)
    ws.row_dimensions[r].height=ht

def _link(ws,r,url):
    if url:
        c=ws.cell(r,1,"Open in Raven"); c.font=_f(10,True,"2563EB"); c.hyperlink=url


def _package(wb,data):
    ws=wb.active; ws.title="Covenant Package"; ws.sheet_view.showGridLines=False
    for col,w in {"A":26,"B":13,"C":24,"D":16,"E":22,"F":30,"G":20}.items(): ws.column_dimensions[col].width=w
    _title(ws,1,data["title"],data.get("meta",""),7); r=4
    if data.get("stale_note"):
        _para(ws,r,"⚠️ "+data["stale_note"],7,ht=30,font=_f(9,color=AMBER_T,italic=True)); r+=1
    if data.get("ebitda_basis"):
        _para(ws,r,"Covenant EBITDA basis: "+data["ebitda_basis"],7,ht=30,font=_f(9,color=GREY,italic=True)); r+=1
    if data.get("covlite_warning"):
        _para(ws,r,"⚠️ "+data["covlite_warning"],7,ht=32,font=_f(10,True,AMBER_T),fillc=AMBER_F); r+=1
    r+=1
    _banner(ws,r,"COVENANT PACKAGE — BY TRANCHE",7); r+=1
    cols=["Facility","Regime","Governed metric","Level","Frequency","Springing trigger","Equity cure"]
    for i,h in enumerate(cols,1):
        c=ws.cell(r,i,h); c.font=_f(10,True,"FFFFFF"); c.fill=_fill(NAVY)
        c.alignment=Alignment(horizontal="left",vertical="center",wrap_text=True,indent=1)
    ws.row_dimensions[r].height=28; r+=1
    for row in data.get("package",[]):
        vals=[row.get("facility"),row.get("regime"),row.get("metric"),row.get("level"),
              row.get("frequency"),row.get("trigger"),row.get("cure")]
        covlite = (row.get("regime","").lower()=="cov-lite")
        for i,v in enumerate(vals,1):
            c=ws.cell(r,i,v if v is not None else "—"); c.border=BORDER
            c.font=_f(9,color=(GREY if covlite else NAVY)); c.alignment=LEFTW if i in (1,3,6) else CENTER
        ws.row_dimensions[r].height=30; r+=1
    if data.get("package_note"):
        _para(ws,r,data["package_note"],7,ht=30,font=_f(9,color=GREY,italic=True)); r+=1
    r+=1
    sm=data.get("summary") or {}
    if sm:
        _banner(ws,r,"SUMMARY",7); r+=1
        for lab,key in [("Compliance","compliance"),("Key risk","key_risk"),("Springing exposure","springing")]:
            if sm.get(key):
                ws.cell(r,1,lab).font=_f(10,True,GREY); ws.cell(r,1).alignment=LEFT
                _para_c=ws.cell(r,2,sm[key]); _para_c.font=_f(10); _para_c.alignment=LEFTW
                ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=7)
                for ci in range(1,8): ws.cell(r,ci).border=BORDER
                ws.row_dimensions[r].height=30; r+=1
    r+=1; _link(ws,r,data.get("web_url"))


def _spread(wb,data):
    spreads=data.get("spreads") or []
    if not spreads: return
    ws=wb.create_sheet("Spread"); ws.sheet_view.showGridLines=False
    for col,w in {"A":16,"B":22,"C":16,"D":12,"E":14}.items(): ws.column_dimensions[col].width=w
    _title(ws,1,data["title"].replace("Covenant Test —","Covenant Spread —"),
           "Each binding covenant spread against the model projections",5); r=4
    for blk in spreads:
        _banner(ws,r,f"{blk.get('facility','')} — {blk.get('metric','')}",5); r+=1
        if blk.get("note"):
            _para(ws,r,blk["note"],5,ht=28,font=_f(9,color=GREY,italic=True)); r+=1
        for i,h in enumerate(["Period","Metric","Threshold","Pass/Fail","Headroom"],1):
            c=ws.cell(r,i,h); c.font=_f(10,True,"FFFFFF"); c.fill=_fill(NAVY); c.alignment=CENTER
        ws.row_dimensions[r].height=22; r+=1
        for period,metric,thr,status,head in blk.get("rows",[]):
            ws.cell(r,1,period).font=_f(10); ws.cell(r,1).alignment=CENTER
            ws.cell(r,2,metric).font=_f(10); ws.cell(r,2).alignment=CENTER
            ws.cell(r,3,thr).font=_f(10); ws.cell(r,3).alignment=CENTER
            c=ws.cell(r,4,status); sf,stc=_status_colors(status); c.fill=_fill(sf); c.font=_f(9,color=stc); c.alignment=CENTER
            ws.cell(r,5,head).font=_f(10); ws.cell(r,5).alignment=CENTER
            for ci in range(1,6): ws.cell(r,ci).border=BORDER
            ws.row_dimensions[r].height=18; r+=1
        r+=1
    if data.get("covlite_note"):
        _para(ws,r,data["covlite_note"],5,ht=34,font=_f(10,color=NAVY)); r+=2
    _link(ws,r,data.get("web_url"))


def _stress(wb,data):
    stress=data.get("stress") or []
    if not stress: return
    ws=wb.create_sheet("Stress"); ws.sheet_view.showGridLines=False
    for col,w in {"A":10,"B":13,"C":11,"D":14,"E":11,"F":11,"G":11,"H":24}.items(): ws.column_dimensions[col].width=w
    _title(ws,1,data["title"].replace("Covenant Test —","Breakeven & Stress —"),
           "EBITDA breakeven and −10 / −20 / −30% stress, per binding covenant",8); r=4
    for blk in stress:
        thr=blk.get("threshold")
        _banner(ws,r,f"{blk.get('facility','')} — {blk.get('metric','')}  (threshold {blk.get('threshold')}x)"
                if isinstance(thr,(int,float)) else f"{blk.get('facility','')} — {blk.get('metric','')}",8); r+=1
        for i,h in enumerate(["Year","Adj. EBITDA","Debt","Base Leverage","−10%","−20%","−30%","Breakeven Decline to Trip"],1):
            c=ws.cell(r,i,h); c.font=_f(10,True,"FFFFFF"); c.fill=_fill(NAVY); c.alignment=CENTER
        ws.row_dimensions[r].height=28; r+=1
        for year,ebitda,debt,base,m10,m20,m30,be in blk.get("rows",[]):
            ws.cell(r,1,year).font=_f(10,True); ws.cell(r,1).alignment=CENTER
            ws.cell(r,2,ebitda).font=_f(10); ws.cell(r,2).alignment=CENTER
            ws.cell(r,3,debt).font=_f(10); ws.cell(r,3).alignment=CENTER
            c=ws.cell(r,4,base); c.number_format='0.00"x"'; c.font=_f(10); c.alignment=CENTER
            for ci,val in ((5,m10),(6,m20),(7,m30)):
                c=ws.cell(r,ci,val); c.number_format='0.00"x"'; c.alignment=CENTER
                lf,lt=_lev_colors(val,thr)
                if lf: c.fill=_fill(lf); c.font=_f(9,color=lt)
                else:  c.font=_f(10)
            ws.cell(r,8,be).font=_f(10); ws.cell(r,8).alignment=CENTER
            for ci in range(1,9): ws.cell(r,ci).border=BORDER
            ws.row_dimensions[r].height=18; r+=1
        r+=1
    _link(ws,r,data.get("web_url"))


def build(data, out_path):
    wb=openpyxl.Workbook()
    _package(wb,data)
    _spread(wb,data)   # no-op if no binding covenants to spread
    _stress(wb,data)   # no-op if no stress blocks
    wb.save(out_path)
    return out_path


if __name__=="__main__":
    demo={
      "title":"Covenant Test — Demo Deal","meta":"USD in millions · Generated today",
      "ebitda_basis":"Adjusted EBITDA as the model determines it.",
      "covlite_warning":"Term Loan has no maintenance covenant; springing revolver test only.",
      "package":[{"facility":"Revolving Facility","regime":"Springing","metric":"First Lien Net Leverage",
                  "level":"9.77x [Assumed]","frequency":"Quarterly, if sprung",
                  "trigger":"Utilization > 40%","cure":"Yes — max 5"},
                 {"facility":"Term Loan","regime":"Cov-lite","metric":"None","level":"—",
                  "frequency":"—","trigger":"—","cure":"—"}],
      "summary":{"compliance":"Pass throughout, as-if-sprung","key_risk":"Cov-lite term loan","springing":"Revolver >40%"},
      "spreads":[{"facility":"Revolving Facility","metric":"First Lien Net Leverage","note":"as-if-sprung",
                  "rows":[("2025E","5.06x","9.77x","Pass","4.71x"),("2026E","4.30x","9.77x","Pass","5.47x")]}],
      "stress":[{"facility":"Revolving Facility","metric":"First Lien Net Leverage","threshold":9.77,
                 "rows":[("2025E","$63.1","$319.3",5.06,5.62,6.32,7.23,"-48.2%")]}],
      "web_url":"https://example.com",
    }
    print("built:", build(demo,"/tmp/_covenant_selftest.xlsx"))
