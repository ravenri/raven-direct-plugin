#!/usr/bin/env python3
"""Clean Raven's Word export of a Credit Assessment.

scan  <export.docx> <work.docx>
      Fixes mechanical defects (misplaced citation markers, split words, em dashes,
      footer wording, table widths, the page break before Assessment Lenses), adds a
      Sources line under each lens heading, splits dense lens narratives into shorter
      paragraphs, saves work.docx, then prints a QA report, the document text
      (so it is read once) and every sentence that carries banned phrasing.

apply <work.docx> <out.docx> [--fixes fixes.tsv] [--notes notes.txt] [--drop "Section heading" ...]
      Applies sentence-level replacements (old<TAB>new per line; digits must not
      change) and appends an Analyst Notes section (one note per line,
      "Label: text", **bold** allowed).
"""
import re, sys, copy, argparse
from collections import Counter
import docx

MARK = re.compile(r'\[(\d+(?:\s*,\s*\d+)*)\]')
SRC_LINE = re.compile(r'^\[\d+\]\s')
BOUND = re.compile(r'[.!?]\s+(?:\[\d+(?:\s*,\s*\d+)*\]\s+)?')
BANNED = re.compile(
    r"credit is sound|likely repaid|supported by|reinforced by|translates? to|more durability|"
    r"\breal\b(?! estate)|robust|compelling|attractive|well-positioned|best-in-class|mission-critical|"
    r"market-leading|world-class|\bpremier\b|stellar|corroborat|\btension\b|arguably|\bquite\b|\bvery\b|"
    r"proven track record|comfortable|\bimpaired\b|\bappropriate\b|\bsound\b|\bstrong\b|\bsignificant\b|"
    r"as structured|most credit-relevant|most material|throws off cash|leans? forward|deleverages? hard|punch line|the story",
    re.I)
PLUMBING = re.compile(r"\bvault\b|\btile\b|prompt id|\bQ\d{3}\b|AI summary", re.I)
NOTE_LABELS = ("Analyst Note", "Key Credit Concern", "Counterargument", "Diligence Item")

def all_paragraphs(d):
    for p in d.paragraphs: yield p
    for t in d.tables:
        for row in t.rows:
            seen = set()
            for c in row.cells:
                if id(c._tc) in seen: continue
                seen.add(id(c._tc))
                for p in c.paragraphs: yield p
    for s in d.sections:
        for part in (s.header, s.footer):
            for p in part.paragraphs: yield p

def ptext(p): return "".join(r.text for r in p.runs)

def edit(p, s, e, new):
    """Replace chars [s,e) of the paragraph text with new, keeping run formatting."""
    runs = p.runs
    if s == e:
        pos = 0
        for i, r in enumerate(runs):
            L = len(r.text)
            if pos < s <= pos + L or (s == 0 and i == 0):
                k = s - pos
                r.text = r.text[:k] + new + r.text[k:]
                return
            pos += L
        runs[-1].text += new
        return
    pos, done = 0, False
    for r in runs:
        t = r.text; rs, re_ = pos, pos + len(t); pos = re_
        if re_ <= s or rs >= e: continue
        a, b = max(s, rs) - rs, min(e, re_) - rs
        if not done: r.text = t[:a] + new + t[b:]; done = True
        else: r.text = t[:a] + t[b:]

def vocab(d):
    c = Counter()
    for p in all_paragraphs(d):
        c.update(w.lower() for w in re.findall(r"[A-Za-z][A-Za-z'\-]*", ptext(p)))
    return c

def fix_markers(d, stats):
    V = vocab(d)
    for p in all_paragraphs(d):
        for _ in range(60):
            text = ptext(p)
            if SRC_LINE.match(text) or re.match(r'^\d+\.\s', text): break
            hit = None
            for m in MARK.finditer(text):
                before = text[:m.start()]
                after = text[m.end():]
                if not re.match(r'\s*[A-Za-z0-9]', after): continue
                if re.search(r'[.!?]\s*$', before): continue
                bs = 0; grp = None
                for bm in BOUND.finditer(before): bs = bm.end(); grp = bm
                seg = before[bs:]
                if not seg.strip() or len(seg) >= 45: continue
                hit = (m, before, after, bs, grp); break
            if not hit: break
            m, before, after, bs, grp = hit
            L = len(before.rstrip()); sp_after = len(after) - len(after.lstrip())
            R = m.end() + sp_after
            left = re.findall(r"[A-Za-z][A-Za-z'\-]*$", before.rstrip())
            join = (sp_after == 0 and left and "-" not in left[0] and left[0].lower() != left[0].upper()
                    and V[left[0].lower()] <= 1 and after[:1].islower())
            if join: stats["joined"] += 1
            nums = m.group(1)
            edit(p, L, R, "" if join else " ")
            if bs == 0:
                edit(p, len(ptext(p)), len(ptext(p)), " [" + nums + "]")
            else:
                prev = re.search(r'(\[(\d+(?:\s*,\s*\d+)*)\])(\s+)$', text[:bs])
                if prev:
                    merged = sorted({int(x) for x in re.findall(r'\d+', prev.group(2) + "," + nums)})
                    edit(p, prev.start(1), prev.end(1), "[" + ", ".join(map(str, merged)) + "]")
                elif "\n" in text[bs - 1 - (len(grp.group()) - 1):bs] if grp else False:
                    edit(p, grp.start() + 1, grp.start() + 1, " [" + nums + "]")
                else:
                    edit(p, bs, bs, "[" + nums + "] ")
            stats["markers"] += 1

def fix_text(d, stats):
    subs = [(" \u2014 ", ", "), ("\u2014", ","), (" \u00b7 \u00b7 ", " \u00b7 "),
            ("deal vault analysis + Financial Analysis", "deal materials")]
    for p in all_paragraphs(d):
        for old, new in subs:
            while old in ptext(p):
                i = ptext(p).index(old); edit(p, i, i + len(old), new)
                if "\u2014" in old: stats["dashes"] += 1

def fix_widths(d, stats):
    """The export sets equal column widths; give each table sensible ones."""
    from docx.shared import Twips
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    for t in d.tables:
        n = len(t.columns); head = t.rows[0].cells[0].text.strip()
        if n == 4: w = [500, 1800, 1750, 5742]
        elif n == 8: w = [3000] + [970] * 7
        elif n == 3 and head == "Metric": w = [2800, 2200, 4792]
        elif n == 3 and head == "Line item": w = [2400, 6200, 1192]
        else: continue
        for g, x in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), w): g.set(qn("w:w"), str(x))
        for row in t.rows:
            for c, x in zip(row.cells, w): c.width = Twips(x)
        pr = t._tbl.tblPr
        lay = pr.find(qn("w:tblLayout"))
        if lay is None: lay = OxmlElement("w:tblLayout"); pr.append(lay)
        lay.set(qn("w:type"), "fixed")
        tw = pr.find(qn("w:tblW"))
        if tw is not None: tw.set(qn("w:w"), str(sum(w))); tw.set(qn("w:type"), "dxa")
        stats["tables"] += 1

CUE = re.compile(r"^(The governing constraint|The binding constraint|The offset|That (?:support |is )?(?:is )?(?:partly )?offset|"
                 r"There are also|That said|However|On the other hand|The constraint)")
SENT = re.compile(r"([.!?](?:\s+\[\d+(?:\s*,\s*\d+)*\])?)(\s+)(?=[A-Z0-9\[(])")

def break_paragraphs(d, stats):
    """Split dense lens narratives into short paragraphs at natural points."""
    paras = d.paragraphs
    heads = [i for i, p in enumerate(paras) if re.match(r'^\d+\.\s', ptext(p)) and MARK.search(ptext(p))]
    for i in heads:
        if i + 1 >= len(paras): continue
        p = paras[i + 1]; text = ptext(p)
        if len(text) < 500: continue
        # sentence spans: (start, end_of_text, ws_start, ws_end)
        sents, last = [], 0
        for m in SENT.finditer(text):
            sents.append((last, m.end(1), m.start(2), m.end(2)))
            last = m.end(2)
        sents.append((last, len(text), None, None))
        cuts, n, ln = [], 0, 0
        for k, (a, b, ws, we) in enumerate(sents[:-1]):
            n += 1; ln += b - a
            if "\n\n" in text[ws:we]: n = ln = 0; continue
            remaining = 0
            for (a2, b2, ws2, we2) in sents[k + 1:]:
                remaining += 1
                if ws2 is not None and "\n\n" in text[ws2:we2]: break
            nxt = text[sents[k + 1][0]:sents[k + 1][1]]
            want = (n >= 2 and CUE.match(nxt)) or n >= 3 or (n >= 2 and ln >= 520)
            if want and (remaining >= 2 or CUE.match(nxt)):
                cuts.append((ws, we)); n = ln = 0
        for ws, we in reversed(cuts):
            edit(p, ws, we, "\n\n"); stats["breaks"] += 1

def add_lens_sources(d, stats):
    """Under each lens, list the documents its heading numbers point to."""
    from docx.text.paragraph import Paragraph
    from docx.shared import Pt, RGBColor
    src = {}
    for p in d.paragraphs:
        m = re.match(r'^\[(\d+)\]\s+(.*)$', ptext(p))
        if m:
            parts = [x.strip() for x in re.split(r'\s+\u00b7\s+', m.group(2))]
            parts[0] = re.sub(r'\.pdf$', '', parts[0], flags=re.I)
            src[int(m.group(1))] = ", ".join(parts)
    paras = d.paragraphs
    heads = [i for i, p in enumerate(paras)
             if re.match(r'^\d+\.\s', ptext(p)) and MARK.search(ptext(p))]
    for i in reversed(heads):
        nums = [int(x) for x in re.findall(r'\d+', MARK.search(ptext(paras[i])).group(1))]
        items = [f"[{n}] {src[n]}" for n in nums if n in src]
        if not items or i + 1 >= len(paras): continue
        body = paras[i + 1]
        el = copy.deepcopy(body._p)
        for r in list(el):
            if r.tag.endswith('}r'): el.remove(r)
        paras[i]._p.addnext(el)
        np_ = Paragraph(el, body._parent)
        hm = MARK.search(ptext(paras[i]))
        edit(paras[i], hm.start() - 1, hm.end(), "")
        from docx.text.run import Run
        tpl = next(r for r in body.runs if not r.bold)._r
        runs = []
        for text, bold in (("Sources  ", True), ("; ".join(items), False)):
            rel = copy.deepcopy(tpl); el.append(rel)
            run = Run(rel, np_); run.text = text; run.bold = bold
            run.font.size = Pt(8); run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        stats["lens_sources"] += 1

def remove_lens_page_break(d, stats):
    """The export forces a page break before Assessment Lenses, leaving a near-empty page."""
    from docx.oxml.ns import qn
    paras = d.paragraphs
    for i, p in enumerate(paras):
        if ptext(p).strip() == "Assessment Lenses" and i > 0:
            prev = paras[i - 1]
            for br in list(prev._p.iter(qn("w:br"))):
                if br.get(qn("w:type")) == "page":
                    br.getparent().remove(br); stats["pagebreaks"] += 1
            return

def sentences(text):
    parts = re.split(r'(?<=[.!?])\s+(?:\[[\d, ]+\]\s+)?', text)
    return [s for s in parts if s.strip()]

def composite_check(d):
    t = d.tables[0]; scores = []
    for r in t.rows[1:]:
        m = re.match(r'\s*(\d)', r.cells[2].text)
        if m: scores.append(int(m.group(1)))
    stated = None
    for p in d.paragraphs:
        m = re.search(r'Overall:\s*([\d.]+)\s*/', p.text)
        if m: stated = m.group(1); break
    if scores: return stated, round(sum(scores) / len(scores), 3), len(scores)
    return stated, None, 0

def scan(src, work):
    d = docx.Document(src)
    stats = Counter()
    fix_markers(d, stats); fix_text(d, stats); fix_widths(d, stats); remove_lens_page_break(d, stats); break_paragraphs(d, stats); add_lens_sources(d, stats)
    d.save(work)
    d = docx.Document(work)
    print("== QA ==")
    print(f"markers moved: {stats['markers']}; words rejoined: {stats['joined']}; em dashes replaced: {stats['dashes']}; tables resized: {stats['tables']}; lens source lines: {stats['lens_sources']}; paragraph breaks added: {stats['breaks']}; page breaks removed: {stats['pagebreaks']}")
    st, comp, n = composite_check(d)
    print(f"composite stated {st}; mean of {n} rated scores {comp}")
    left = [ptext(p)[:90] for p in all_paragraphs(d) if "\u2014" in ptext(p)]
    print("em dashes remaining:", len(left))
    pl = [(ptext(p)[:90]) for p in all_paragraphs(d) if PLUMBING.search(ptext(p))]
    print("platform wording found:", pl if pl else "none")
    print("\n== DOCUMENT TEXT ==")
    for i, p in enumerate(d.paragraphs):
        if ptext(p).strip() == "Sources": break
        if ptext(p).strip(): print(f"P{i}: {ptext(p)}")
    for ti, t in enumerate(d.tables):
        if ti == 0: continue
        for ri, row in enumerate(t.rows):
            print(f"T{ti}r{ri}: " + " | ".join(c.text.replace("\n", " ") for c in row.cells))
    print("\n== BANNED PHRASING (sentence-level) ==")
    k = 0
    for p in all_paragraphs(d):
        if SRC_LINE.match(ptext(p)): continue
        for s in sentences(ptext(p)):
            if BANNED.search(s):
                k += 1; print(f"- {s.strip()}")
    print(f"({k} sentences)")

def add_notes(d, notes):
    paras = d.paragraphs
    idx = next(i for i, p in enumerate(paras) if p.text.startswith("Generated by"))
    anchor = paras[idx]._p
    h_t = next(p for p in paras if p.text.strip() == "Cross-Lens Checks")
    lab_t = next(p for p in paras if p.text.startswith("Earnings quality / "))
    body_t = paras[paras.index(lab_t) + 1]
    plain = next(r for r in body_t.runs if not r.bold)._r
    bold = next(r for r in body_t.runs if r.bold)._r
    def clone(tpl_p, runs_spec, own=False):
        plain_r = tpl_p.runs[0]._r if own else plain
        bold_r = tpl_p.runs[0]._r if own else bold
        el = copy.deepcopy(tpl_p._p)
        for r in list(el):
            if r.tag.endswith('}r'): el.remove(r)
        for txt, is_bold in runs_spec:
            r = copy.deepcopy(bold_r if is_bold else plain_r); 
            for t in r.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'): t.text = txt; t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            el.append(r)
        anchor.addprevious(el)
    clone(h_t, [("Analyst Notes", True)], own=True)
    for line in notes:
        if ":" not in line: continue
        label, text = line.split(":", 1)
        label, text = label.strip(), text.strip()
        if label not in NOTE_LABELS: raise SystemExit(f"Unknown note label: {label}")
        clone(lab_t, [(label, True)], own=True)
        spec = []
        for i, seg in enumerate(re.split(r'\*\*', text)):
            if seg: spec.append((seg, i % 2 == 1))
        clone(body_t, spec)

def drop_sections(d, names):
    """Delete a sub-section: its heading, the table(s) that follow, and empty spacers."""
    for name in names:
        for p in d.paragraphs:
            if ptext(p).strip().lower() == name.strip().lower():
                el = p._p; nxt = el.getnext(); el.getparent().remove(el)
                while nxt is not None:
                    tag = nxt.tag.split("}")[1]
                    if tag == "tbl" or (tag == "p" and not "".join(nxt.itertext()).strip()):
                        after = nxt.getnext(); nxt.getparent().remove(nxt); nxt = after
                    else: break
                print("dropped section:", name); break
        else: print("section NOT FOUND:", name)

def apply(work, out, fixes, notes, drop=None):
    d = docx.Document(work)
    if drop: drop_sections(d, drop)
    if fixes:
        for ln in open(fixes, encoding="utf-8").read().splitlines():
            if "\t" not in ln: continue
            old, new = ln.split("\t", 1)
            if re.findall(r'\d+', old) != re.findall(r'\d+', new):
                print("REFUSED (digits differ):", old[:60]); continue
            n = 0
            for p in all_paragraphs(d):
                while old in ptext(p):
                    i = ptext(p).index(old); edit(p, i, i + len(old), new); n += 1
            print(("applied " if n else "NOT FOUND ") + old[:60])
    if notes:
        add_notes(d, [l for l in open(notes, encoding="utf-8").read().splitlines() if l.strip()])
    d.save(out)
    left = sum(len(BANNED.findall(s)) for p in all_paragraphs(d) if not SRC_LINE.match(ptext(p)) for s in sentences(ptext(p)))
    print(f"saved {out}; banned phrases still present: {left}")

if __name__ == "__main__":
    ap = argparse.ArgumentParser(); sub = ap.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("scan"); a.add_argument("src"); a.add_argument("work")
    b = sub.add_parser("apply"); b.add_argument("work"); b.add_argument("out")
    b.add_argument("--fixes"); b.add_argument("--notes"); b.add_argument("--drop", nargs="*")
    n = ap.parse_args()
    scan(n.src, n.work) if n.cmd == "scan" else apply(n.work, n.out, n.fixes, n.notes, n.drop)
