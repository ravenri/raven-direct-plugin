#!/usr/bin/env python3
"""Scan a finished .docx for house-rule violations. Exit 1 if any are found."""
import re, sys
from docx import Document

BANNED = [
    r"\bthe vault\b", r"\bvault\b", r"\btiles?\b(?!\w)", r"\bdashboards?\b", r"\bQ\d{3,4}\b", r"diligence_summary",
    r"\bbest-in-class\b", r"\brobust\b", r"\bcompelling\b", r"\battractive\b", r"\bwell-positioned\b",
    r"\bmission-critical\b", r"\bworld-class\b", r"\bmarket-leading\b", r"\bpremier\b", r"\bstellar\b",
    r"\bproven track record\b", r"\bcorroborated\b", r"\btranslates to\b", r"\btension\b", r"\barguably\b",
    r"\bthe story\b", r"\bpoised to\b", r"\breal (cushion|concentration|risk)\b",
    r"\bsignificant(ly)?\b(?!\s+\d)",
]
BRITISH = r"\b(favour\w*|colour\w*|behaviour\w*|organis\w+|analys(e|ed|es|ing)|centre\w*|licence\w*|\w+isation\w*)\b"
BAD_HEADINGS = {"overview", "assessment", "the read", "the take", "summary"}

d = Document(sys.argv[1])
parts = []
for p in d.paragraphs:
    parts.append((p.style.name if p.style is not None else "", p.text))
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            parts.append(("cell", c.text))
text = "\n".join(t for _, t in parts)
problems = []
for ch, name in (("\u2014", "em dash"), ("\u2013", "en dash")):
    if ch in text: problems.append(f"{name} x{text.count(ch)}")
for pat in BANNED:
    for m in re.finditer(pat, text, flags=re.I):
        problems.append(f"banned: '{m.group(0)}'")
for m in re.finditer(BRITISH, text, flags=re.I):
    problems.append(f"British spelling: '{m.group(0)}'")
for style, t in parts:
    if style.startswith("Heading") and re.sub(r"^\d+\.\s*", "", t.strip().lower()) in BAD_HEADINGS:
        problems.append(f"function-label heading: '{t}'")
if problems:
    print("CHECK FAILED"); [print(" -", x) for x in sorted(set(problems))]; sys.exit(1)
print("CHECK PASSED")
