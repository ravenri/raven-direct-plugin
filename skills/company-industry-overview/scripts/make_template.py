"""One-time builder for assets/industry-overview-template.docx (layout and style guideline).
Run only when the house style changes. Contains no deal content."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import parse_xml, OxmlElement

NAVY = RGBColor(0x1F, 0x38, 0x64)
GRAY = RGBColor(0x59, 0x59, 0x59)
import os, subprocess
if not os.path.exists("assets/base.docx"):
    open("assets/base.docx","wb").write(subprocess.check_output(["pandoc","--print-default-data-file","reference.docx"]))
d = Document("assets/base.docx")

def font(style, size, bold=None, color=None, italic=None):
    f = style.font
    f.name = "Arial"; f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color is not None: f.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), "Arial")
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]

S = d.styles
BY = {st.name: st for st in S}
BY.update({st.style_id: st for st in S})
def G(n):
    return BY.get(n) or BY[n.replace(" ","")]
for n in ("Normal", "Body Text", "First Paragraph", "Compact"):
    font(G(n), 10)
for n in ("Body Text", "First Paragraph"):
    pf = G(n).paragraph_format; pf.space_after = Pt(6); pf.space_before = Pt(0); pf.line_spacing = 1.15
G("Compact").paragraph_format.space_after = Pt(3)
font(G("Title"), 20, True, NAVY); G("Title").paragraph_format.space_after = Pt(4)
for n, sz, sb, sa in (("Heading 1", 13, 16, 6), ("Heading 2", 11, 11, 4), ("Heading 3", 10, 7, 2)):
    font(G(n), sz, True, NAVY if n != "Heading 3" else RGBColor(0, 0, 0))
    G(n).paragraph_format.space_before = Pt(sb); G(n).paragraph_format.space_after = Pt(sa)
    G(n).paragraph_format.keep_with_next = True
# Heading 1 bottom rule
p = G("Heading 1").element.get_or_add_pPr()
p.append(parse_xml('<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                   '<w:bottom w:val="single" w:sz="6" w:space="2" w:color="1F3864"/></w:pBdr>'))

def custom(name, size, bold=False, italic=False, color=None, before=0, after=6, bold_lead=False):
    st = S.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = G("Body Text")
    font(st, size, bold, color, italic)
    st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)
    return st
custom("Meta", 9.5, after=4)
custom("Source Line", 8, italic=True, color=GRAY, after=8)
custom("Note", 8.5, italic=True, color=GRAY, after=6)
custom("Credit Implication", 10, before=3, after=8)

# Table style: light gray header row, thin gray borders, 8.5 pt text
ts = G("Table")
font(ts, 8.5)
tblPr = ts.element.find(qn("w:tblPr"))
if tblPr is None:
    tblPr = OxmlElement("w:tblPr"); ts.element.append(tblPr)
for ch in list(tblPr):
    if ch.tag in (qn("w:tblBorders"), qn("w:tblCellMar")): tblPr.remove(ch)
W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
tblPr.append(parse_xml(f'<w:tblBorders {W}>' + "".join(
    f'<w:{s} w:val="single" w:sz="4" w:space="0" w:color="BFBFBF"/>'
    for s in ("top", "left", "bottom", "right", "insideH", "insideV")) + '</w:tblBorders>'))
tblPr.append(parse_xml(f'<w:tblCellMar {W}><w:top w:w="50" w:type="dxa"/><w:left w:w="90" w:type="dxa"/>'
                       '<w:bottom w:w="50" w:type="dxa"/><w:right w:w="90" w:type="dxa"/></w:tblCellMar>'))
for old in ts.element.findall(qn("w:tblStylePr")): ts.element.remove(old)
ts.element.append(parse_xml(f'<w:tblStylePr {W} w:type="firstRow"><w:rPr><w:b/></w:rPr>'
                            '<w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="E7E9EE"/></w:tcPr></w:tblStylePr>'))

# Page setup and footer
sec = d.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.left_margin = sec.right_margin = Inches(1); sec.top_margin = sec.bottom_margin = Inches(0.85)
fp = sec.footer.paragraphs[0]
fp.alignment = 1
r = fp.add_run("Page "); r.font.size = Pt(8); r.font.name = "Arial"; r.font.color.rgb = GRAY
r2 = fp.add_run(); r2.font.size = Pt(8); r2.font.name = "Arial"; r2.font.color.rgb = GRAY
for t, txt in (("begin", None), (None, " PAGE "), ("end", None)):
    if t:
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t)
    else:
        e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
    r2._r.append(e)
d.save("assets/industry-overview-template.docx")
os.remove("assets/base.docx")
print("template written")
