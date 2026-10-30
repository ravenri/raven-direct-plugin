"""Pre-delivery check for the Deals at a Glance workbook.

Usage: python check_deals_at_a_glance.py INPUT.json RECALCULATED.xlsx
Run after recalculation, with the same input file the workbook was built from.
"""
import json, sys, os, re, math, datetime
from openpyxl import load_workbook
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_deals_at_a_glance as B

BANNED = ["vault", "prompt", "tile", "dashboard", "mcp", "—"]

def main(inp, xlsx):
    cfg = json.load(open(inp))
    gics = B.sectors()
    for d in cfg["deals"]:
        d["_p"] = B.parse_deal(d)[0]
    wb = load_workbook(xlsx, data_only=True)
    ws = wb["Averages by Year"]
    errs, bad = [], []
    for sh in wb:
        for row in sh.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    if c.value.startswith("#"): errs.append(f"{sh.title}!{c.coordinate} {c.value}")
                    low = c.value.lower()
                    for w in BANNED:
                        if (w in low) and not (w == "tile" and "tiles" not in low and "tile " not in low and not re.search(r"\btile\b", low)):
                            bad.append(f"{sh.title}!{c.coordinate} has '{w}'")
    rows = {}
    for r in range(5, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(ws.cell(r, 2).value, datetime.datetime):
            rows[a] = r
    out = []
    n = len(cfg["deals"])
    out.append(("deal rows match input", len(rows) == n))
    out.append(("first of month dates", all(ws.cell(r, 2).value.day == 1 for r in rows.values())))
    out.append(("valid GICS sectors", all(d["gics_sector"] in gics for d in cfg["deals"])))
    def avg(vals):
        v = [x for x in vals if x is not None]
        return sum(v) / len(v) if v else None
    close = lambda a, b: (a in (None, "") and b is None) or (isinstance(a, (int, float)) and b is not None and math.isclose(a, b, rel_tol=1e-6, abs_tol=1e-9))
    ok_all = True
    for j, k in B.KEYS.items():
        exp = avg([d["_p"][k] for d in cfg["deals"]])
        ok_all &= close(next((ws.cell(r, j).value for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value == "All deals average"), None), exp)
        for yr in {int(d["date"][:4]) for d in cfg["deals"]}:
            e = avg([d["_p"][k] for d in cfg["deals"] if int(d["date"][:4]) == yr])
            cell = next((ws.cell(r, j).value for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value == f"{yr} average"), None)
            ok_all &= close(cell, e)
        for s in gics:
            e = avg([d["_p"][k] for d in cfg["deals"] if d["gics_sector"] == s])
            cell = next((ws.cell(r, j).value for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value == s), None)
            ok_all &= close(cell, e)
    out.append(("year, all deals and sector averages tie to inputs", ok_all))
    tot = next((ws.cell(r, 3).value for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value == "Total deals"), None)
    out.append(("sector table totals the deal count", tot == n))
    out.append(("no formula errors", not errs))
    out.append(("no banned wording or em dashes", not bad))
    for name, ok in out: print(("PASS " if ok else "FAIL ") + name)
    for e in errs + bad: print("  " + e)
    sys.exit(0 if all(ok for _, ok in out) else 1)

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
