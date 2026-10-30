"""Build the Deals at a Glance workbook from the template.

Usage: python build_deals_at_a_glance.py INPUT.json OUTPUT.xlsx

INPUT.json is a build input, not a deliverable. Keep it in scratch space.
Shape:
  {"deals": [{"deal": name, "date": "YYYY-MM-01", "gics_sector": one of the 11,
              "currency": "USD", "glance": [[label, value], ...] as shown in Raven}],
   "notes": ["run-specific note", ...]}
Values are parsed here from the text Raven shows. Nothing is corrected or restated.
Year averages, the all deals average and the sector table are live formulas.
"""
import json, re, sys, os, datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "..", "assets", "Deals_at_a_Glance_Template.xlsx")
GICS = os.path.join(HERE, "..", "assets", "GICS_Sectors.xlsx")
thin, med = Side(style="thin"), Side(style="medium")
NAVY, BAND, AVG = "1F3864", "D9E1F2", "F2F2F2"
FMT = {4: '#,##0', 5: '0.0\\x', 6: '0.00\\x', 7: '0.00\\x', 8: '0.0%', 9: '0.0%'}
KEYS = {4: "ev", 5: "ev_ebitda", 6: "net", 7: "total", 8: "ltv", 9: "equity"}
METRICS = "DEFGHI"

def F(b=False, i=False, s=11, c=None):
    return Font(name="Arial", bold=b, italic=i, size=s, color=c)

def put(ws, r, c, v, font=None, fmt=None, al=None, top=None, bottom=None, fill=None, wrap=False):
    x = ws.cell(r, c, v)
    x.font = font or F()
    if fmt: x.number_format = fmt
    if al or wrap: x.alignment = Alignment(horizontal=al, wrap_text=wrap, vertical="top" if wrap else None)
    if top or bottom: x.border = Border(top=top, bottom=bottom)
    if fill: x.fill = PatternFill("solid", fgColor=fill)
    return x

def sectors():
    ws = load_workbook(GICS)["GICS Sectors"]
    return [ws.cell(r, 2).value for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value]

# ---- parsing of the text Raven shows -------------------------------------
MONEY = re.compile(r"(?P<cur>[A-Za-z]{0,3}\$|€|£)\s?(?P<n>[\d,]+(?:\.\d+)?)\s?(?P<u>B|MM|M|K)?\b")

def money(text):
    """Return (value in $mm or None, flag) from the first amount in text."""
    ms = list(MONEY.finditer(text))
    if not ms:
        return None, "none"
    m = ms[0]
    if m.group("cur") not in ("$", "US$"):
        return None, "nonusd"
    n = float(m.group("n").replace(",", ""))
    u = m.group("u")
    if u is None:
        return None, "nounit"
    v = n * 1000 if u == "B" else n / 1000 if u == "K" else n
    return v, ("two" if len(ms) > 1 else "ok")

def mult(text):
    m = re.search(r"(\d+(?:\.\d+)?)\s*x", text)
    return float(m.group(1)) if m else None

def pct(text):
    m = re.search(r"(\d+(?:\.\d+)?)\s*%", text)
    return float(m.group(1)) / 100 if m else None

def parse_deal(d):
    p = {k: None for k in KEYS.values()}
    raw = {"size": "", "netdebt": "", "cash": "", "ev": "", "evx": "", "net": "", "total": "", "ltv": "", "equity": ""}
    other, flags = [], []
    for label, value in d["glance"]:
        l = label.lower()
        if l.startswith("enterprise value"):
            raw["ev"] = value
            v, f = money(value)
            if d.get("currency", "USD").upper() != "USD" or f == "nonusd":
                flags.append("EV left out: not in US dollars")
            else:
                p["ev"] = v
                if f == "two": flags.append("EV two values shown, first used")
                if f == "nounit": flags.append("EV has no unit, left blank")
        elif l.startswith("ev /") or l.startswith("ev/"):
            raw["evx"] = value; p["ev_ebitda"] = mult(value)
        elif l.startswith("net leverage"):
            raw["net"] = value; p["net"] = mult(value)
            if re.search(r"secured|first lien|1l", value + " " + label, re.I):
                flags.append("Net leverage is " + ("secured net" if re.search("secured", value + label, re.I) else "first lien net"))
        elif l.startswith("total leverage"):
            raw["total"] = value; p["total"] = mult(value)
            if re.search(r"\bnet\b", value, re.I) and not re.search("gross", value, re.I):
                flags.append("Total leverage is shown net")
        elif l.startswith("first lien ltv"):
            raw["ltv"] = value; p["ltv"] = pct(value)
        elif l.startswith("sponsor equity") or l.startswith("equity cushion"):
            raw["equity"] = value; p["equity"] = pct(value)
            if l.startswith("equity cushion"):
                flags.append("Equity cushion % shown instead of sponsor equity %")
        elif l.startswith("transaction size"): raw["size"] = value
        elif l.startswith("total net debt"): raw["netdebt"] = value
        elif l.startswith("pro forma cash"): raw["cash"] = value
        else: other.append(f"{label}: {value}")
    return p, raw, "; ".join(other), flags

def main(inp, out):
    cfg = json.load(open(inp))
    gics = sectors()
    deals = sorted(cfg["deals"], key=lambda d: d["date"])
    for d in deals:
        d["_dt"] = datetime.datetime.strptime(d["date"], "%Y-%m-%d")
        if d["gics_sector"] not in gics:
            raise SystemExit(f"{d['deal']}: '{d['gics_sector']}' is not one of the 11 GICS sectors")
        d["_p"], d["_raw"], d["_other"], d["_flags"] = parse_deal(d)
    wb = load_workbook(TEMPLATE)
    ws = wb["Averages by Year"]
    r = 5
    avg_rows = []
    for yr in sorted({d["_dt"].year for d in deals}):
        for j in range(1, 10): put(ws, r, j, str(yr) if j == 1 else None, F(True), fill=BAND)
        r += 1
        first = r
        for d in [x for x in deals if x["_dt"].year == yr]:
            put(ws, r, 1, d["deal"], bottom=thin)
            put(ws, r, 2, d["_dt"], fmt="mmm yyyy", al="center", bottom=thin)
            put(ws, r, 3, d["gics_sector"], bottom=thin)
            for j in range(4, 10):
                v = d["_p"][KEYS[j]]
                put(ws, r, j, v, F(c="0000FF"), fmt=FMT[j], bottom=thin)
            r += 1
        last = r - 1
        put(ws, r, 1, f"{yr} average", F(True), top=thin, fill=AVG)
        put(ws, r, 2, None, top=thin, fill=AVG)
        put(ws, r, 3, f'=COUNTA(A{first}:A{last})&IF(COUNTA(A{first}:A{last})=1," deal"," deals")', F(True), top=thin, fill=AVG)
        for j in range(4, 10):
            c = METRICS[j - 4]
            put(ws, r, j, f'=IFERROR(AVERAGE({c}{first}:{c}{last}),"")', F(True), fmt=FMT[j], top=thin, fill=AVG)
        avg_rows.append(r)
        r += 2
    end = r - 2
    put(ws, r, 1, "All deals average", F(True), top=med)
    put(ws, r, 2, None, top=med)
    put(ws, r, 3, f'=COUNT(B5:B{end})&" deals"', F(True), top=med)
    for j in range(4, 10):
        c = METRICS[j - 4]
        put(ws, r, j, f'=IFERROR(AVERAGEIFS({c}5:{c}{end},$B5:$B{end},">0"),"")', F(True), fmt=FMT[j], top=med)
    r += 2
    put(ws, r, 1, "Averages by GICS sector", F(True, s=12)); r += 1
    for j, h in {1: "GICS Sector", 3: "Deals", 4: "Enterprise value ($mm)", 5: "EV / EBITDA", 6: "Net leverage",
                 7: "Total leverage", 8: "First lien LTV", 9: "Sponsor equity / cushion %"}.items():
        put(ws, r, j, h, F(True, c="FFFFFF"), al="center", fill=NAVY, wrap=True)
    put(ws, r, 2, None, fill=NAVY)
    r += 1
    s0 = r
    for s in gics:
        put(ws, r, 1, s, bottom=thin); put(ws, r, 2, None, bottom=thin)
        put(ws, r, 3, f"=COUNTIF($C$5:$C${end},A{r})", al="center", bottom=thin)
        for j in range(4, 10):
            c = METRICS[j - 4]
            put(ws, r, j, f'=IFERROR(AVERAGEIFS(${c}$5:${c}${end},$C$5:$C${end},$A{r}),"")', fmt=FMT[j], bottom=thin)
        r += 1
    put(ws, r, 1, "Total deals", F(True), top=thin)
    put(ws, r, 3, f"=SUM(C{s0}:C{r-1})", F(True), al="center", top=thin)
    # detail tab
    wd = wb["Deal at a Glance"]
    for i, d in enumerate(deals):
        rr = 5 + i
        raw = d["_raw"]
        vals = [d["deal"], d["_dt"], d["gics_sector"], d.get("currency", ""), raw["size"], raw["ev"], raw["evx"],
                raw["netdebt"], raw["net"], raw["total"], raw["ltv"], raw["cash"], raw["equity"], d["_other"]]
        for j, v in enumerate(vals, 1):
            put(wd, rr, j, v or None, fmt="mmm yyyy" if j == 2 else None, al="center" if j == 2 else None,
                bottom=thin, wrap=(j == 14))
    # notes
    ns = wb["Notes"]
    n = len(deals)
    fixed = [
        f"Scope: {n} deals with status Closed in Raven. Source: the Deal At A Glance section of each deal's summary. Numbers on the first tab are parsed from the values Raven shows, which are kept as shown on the Deal at a Glance tab.",
        "Date: the year and month in the Raven deal name, shown as month and year. A deal's name may not match its IC memo date.",
        "GICS Sector: one of the 11 GICS sectors, assigned from the company's principal business as described in the deal summary.",
        "Averages are simple averages across deals. They skip any deal where Raven shows no value for that metric, so the number of deals behind each average differs by column. Sectors with no deals show a blank.",
        "Leverage bases differ by deal. Net and total leverage are on each deal's own EBITDA and period, as shown on the Deal at a Glance tab. Treat cross-deal averages as indicative.",
        "Figures are taken as Raven shows them. Where Raven shows two values for one metric, the first is used. Enterprise value is left out where Raven shows it in a currency other than US dollars.",
    ]
    auto = []
    def deals_with(txt): return [d["deal"] for d in deals if txt in d["_flags"]]
    for key, msg in [("Net leverage is secured net", "Net leverage is on a secured net basis for"),
                     ("Net leverage is first lien net", "Net leverage is on a first lien net basis for"),
                     ("Total leverage is shown net", "Raven shows total leverage on a net basis for"),
                     ("Equity cushion % shown instead of sponsor equity %", "Equity cushion % is shown in place of sponsor equity % for"),
                     ("EV left out: not in US dollars", "Enterprise value is left out, because it is not in US dollars, for"),
                     ("EV two values shown, first used", "Raven shows two enterprise values, and the first is used, for")]:
        x = deals_with(key)
        if x: auto.append(f"{msg}: {', '.join(x)}.")
    for label, k in [("enterprise value", "ev"), ("EV / EBITDA", "ev_ebitda"), ("net leverage", "net"),
                     ("total leverage", "total"), ("first lien LTV", "ltv"), ("sponsor equity %", "equity")]:
        x = [d["deal"] for d in deals if d["_p"][k] is None and not (k == "ev" and d["_raw"]["ev"])]
        if x: auto.append(f"No {label} shown in Raven for: {', '.join(x)}.")
    for i, t in enumerate(fixed + auto + cfg.get("notes", [])):
        c = ns.cell(3 + i, 1, t); c.font = F(s=10); c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(out)
    print(f"wrote {out}: {n} deals")

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
