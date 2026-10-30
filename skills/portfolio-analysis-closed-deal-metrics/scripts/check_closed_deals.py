"""Pre-delivery check for the Closed Deals by Year workbook.

Usage: python check_closed_deals.py WORKBOOK.xlsx INPUT.json
Run after recalculation. Compares the live formula results with counts taken from the build input.
Exits 1 on failure.
"""
import sys, re, json, os
from openpyxl import load_workbook

BANNED = ["vault", "prompt", "tile", "dashboard", "mcp"]
GICS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "assets", "GICS_Sectors.xlsx")

def main(path, inp):
    cfg = json.load(open(inp)); deals = cfg["deals"]
    g = load_workbook(GICS)["GICS Sectors"]
    gics = {g.cell(r, 2).value for r in range(5, g.max_row + 1) if g.cell(r, 2).value}
    wv = load_workbook(path, data_only=True)
    ws = wv["Closed Deals"]
    bad = []
    n = len(deals)
    amt = sum(d["amount"] or 0 for d in deals)
    rows = {str(c.value): c.row for c in ws["A"] if c.value}
    ay = rows.get("All years")
    if ay is None: bad.append("No All years row")
    elif ws.cell(ay, 6).value != amt: bad.append(f"All years commitments {ws.cell(ay,6).value} vs input {amt}")
    for yr in sorted({d["date"][:4] for d in deals}):
        want_n = sum(1 for d in deals if d["date"][:4] == yr)
        want_a = sum(d["amount"] or 0 for d in deals if d["date"][:4] == yr)
        r = rows.get(f"{yr} total")
        if r is None: bad.append(f"Missing {yr} total"); continue
        if ws.cell(r, 6).value != want_a: bad.append(f"{yr} total commitments {ws.cell(r,6).value} vs {want_a}")
        if not str(ws.cell(r, 5).value).startswith(str(want_n)): bad.append(f"{yr} total count {ws.cell(r,5).value} vs {want_n}")
    memo = rows.get("Memo: closed deals")
    if memo is None or ws.cell(memo, 2).value != n: bad.append("Memo deal count does not match the number of deals")
    tot = rows.get("Total")
    if tot is None: bad.append("No Total row under By GICS sector")
    else:
        if ws.cell(tot, 2).value != n: bad.append(f"GICS sector table counts {ws.cell(tot,2).value}, expected {n}")
        if ws.cell(tot, 3).value != amt: bad.append(f"GICS sector table commitments {ws.cell(tot,3).value}, expected {amt}")
    for d in deals:
        if d["gics_sector"] not in gics: bad.append(f"{d['deal']}: not a GICS sector")
        for k in ("date", "business", "sponsors", "sponsor_label"):
            if not d.get(k): bad.append(f"{d['deal']}: {k} is empty")
        if not re.match(r"^\d{4}-\d{2}-01$", d["date"]): bad.append(f"{d['deal']}: date must be the first of the month")
    if "Total (sponsors overlap)" not in rows: bad.append("No sponsor total row")
    for sh in wv.worksheets:
        for row in sh.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str):
                    if v.startswith("#"): bad.append(f"{sh.title}!{c.coordinate} error {v}")
                    if "—" in v: bad.append(f"{sh.title}!{c.coordinate} has an em dash")
                    for b in BANNED:
                        if re.search(r"\b" + b + r"\b", v.lower()): bad.append(f"{sh.title}!{c.coordinate} contains '{b}'")
    if bad:
        print("FAILED"); [print(" -", x) for x in bad]; sys.exit(1)
    print(f"OK: {n} deals, {amt:g} total commitments")

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
