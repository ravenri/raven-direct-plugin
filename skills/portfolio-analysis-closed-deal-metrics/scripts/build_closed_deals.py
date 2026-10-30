"""Build the Closed Deals by Year workbook from the template.

Usage: python build_closed_deals.py INPUT.json OUTPUT.xlsx

INPUT.json is a build input, not a deliverable. Keep it in scratch space.
Shape:
  {"deals": [{"deal","date" (YYYY-MM-DD, first of the month),"sponsor_label",
              "gics_sector" (one of the 11 in assets/GICS_Sectors.xlsx),
              "business" (short description),"amount" (number or null),
              "sponsors":["Name", ...up to 4]}],
   "notes": ["run-specific note", ...]}
Year blocks, totals and summary tables are live formulas.
"""
import json, sys, os, datetime
from collections import Counter
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "..", "assets", "Closed_Deal_Metrics_Template.xlsx")
GICS = os.path.join(HERE, "..", "assets", "GICS_Sectors.xlsx")
NUM = '#,##0;\\(#,##0\\);\\-'
thin, med = Side(style="thin"), Side(style="medium")
NAVY, BAND = "1F3864", "D9E1F2"
LASTCOL = 10  # column J

def F(b=False, i=False, s=11, c=None):
    return Font(name="Arial", bold=b, italic=i, size=s, color=c)

def put(ws, r, c, v, font=None, fmt=None, al=None, top=None, bottom=None, fill=None):
    x = ws.cell(r, c, v)
    x.font = font or F()
    if fmt: x.number_format = fmt
    if al: x.alignment = Alignment(horizontal=al)
    if top or bottom: x.border = Border(top=top, bottom=bottom)
    if fill: x.fill = PatternFill("solid", fgColor=fill)
    return x

def header(ws, r, labels):
    for j, h in enumerate(labels, 1):
        put(ws, r, j, h, F(True, c="FFFFFF"), al="center", fill=NAVY)

def sectors():
    ws = load_workbook(GICS)["GICS Sectors"]
    return [ws.cell(r, 2).value for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value]

def main(inp, out):
    cfg = json.load(open(inp))
    gics = sectors()
    deals = sorted(cfg["deals"], key=lambda d: d["date"])  # stable within a date
    for d in deals:
        d["_dt"] = datetime.datetime.strptime(d["date"], "%Y-%m-%d")
        if d["gics_sector"] not in gics:
            raise SystemExit(f"{d['deal']}: '{d['gics_sector']}' is not one of the 11 GICS sectors")
    wb = load_workbook(TEMPLATE)
    ws = wb["Closed Deals"]
    r = 5
    year_total_rows = []
    for yr in sorted({d["_dt"].year for d in deals}):
        for j in range(1, LASTCOL + 1):
            put(ws, r, j, str(yr) if j == 1 else None, F(True), fill=BAND)
        r += 1
        first = r
        for d in [x for x in deals if x["_dt"].year == yr]:
            sp = (d.get("sponsors") or [])[:4]
            vals = [d["deal"], d["_dt"], d["sponsor_label"], d["gics_sector"], d["business"], d.get("amount")] + sp
            for j, v in enumerate(vals, 1):
                font = F(c="0000FF") if j == 6 and v is not None else F()
                put(ws, r, j, v, font, fmt={2: "mmm yyyy", 6: NUM}.get(j), al="center" if j == 2 else None, bottom=thin)
            for j in range(len(vals) + 1, LASTCOL + 1):
                put(ws, r, j, None, bottom=thin)
            r += 1
        last = r - 1
        put(ws, r, 1, f"{yr} total", F(True), top=thin)
        put(ws, r, 5, f'=COUNTA(A{first}:A{last})&IF(COUNTA(A{first}:A{last})=1," deal"," deals")', F(True), top=thin)
        put(ws, r, 6, f"=SUM(F{first}:F{last})", F(True), fmt=NUM, top=thin)
        for j in (2, 3, 4, 7, 8, 9, 10): put(ws, r, j, None, top=thin)
        year_total_rows.append(r)
        r += 2
    data_end = r - 1
    put(ws, r, 1, "All years", F(True), top=med)
    put(ws, r, 6, "=" + "+".join(f"F{x}" for x in year_total_rows), F(True), fmt=NUM, top=med)
    for j in (2, 3, 4, 5, 7, 8, 9, 10): put(ws, r, j, None, top=med)
    allyears = r
    rng = lambda col: f"${col}$5:${col}${data_end}"
    amt = rng("F")
    r += 2
    put(ws, r, 1, "By GICS sector", F(True, s=12)); r += 1
    header(ws, r, ["GICS Sector", "Deals", "Commitments ($mm)"]); r += 1
    f0 = r
    for s in gics:
        put(ws, r, 1, s, bottom=thin)
        put(ws, r, 2, f"=COUNTIF({rng('D')},A{r})", al="center", bottom=thin)
        put(ws, r, 3, f"=SUMIF({rng('D')},A{r},{amt})", fmt=NUM, bottom=thin)
        r += 1
    put(ws, r, 1, "Total", F(True), top=thin)
    put(ws, r, 2, f"=SUM(B{f0}:B{r-1})", F(True), al="center", top=thin)
    put(ws, r, 3, f"=SUM(C{f0}:C{r-1})", F(True), fmt=NUM, top=thin)
    r += 2
    cnt = Counter(); order = {}
    for d in deals:
        for s in (d.get("sponsors") or [])[:4]:
            cnt[s] += 1; order.setdefault(s, len(order))
    sponsors = sorted(cnt, key=lambda s: (-cnt[s], order[s]))
    put(ws, r, 1, "By sponsor", F(True, s=12)); r += 1
    header(ws, r, ["Sponsor", "Deals", "Commitments ($mm)"]); r += 1
    s0 = r
    for s in sponsors:
        put(ws, r, 1, s, bottom=thin)
        put(ws, r, 2, "=" + "+".join(f"COUNTIF({rng(c)},A{r})" for c in "GHIJ"), al="center", bottom=thin)
        put(ws, r, 3, "=" + "+".join(f"SUMIF({rng(c)},A{r},{amt})" for c in "GHIJ"), fmt=NUM, bottom=thin)
        r += 1
    put(ws, r, 1, "Total (sponsors overlap)", F(True), top=thin)
    put(ws, r, 2, f"=SUM(B{s0}:B{r-1})", F(True), al="center", top=thin)
    put(ws, r, 3, f"=SUM(C{s0}:C{r-1})", F(True), fmt=NUM, top=thin)
    r += 1
    put(ws, r, 1, "Memo: closed deals", F(i=True))
    put(ws, r, 2, f"=COUNT({rng('B')})", F(i=True), al="center")
    put(ws, r, 3, f"=F{allyears}", F(i=True), fmt=NUM)
    r += 1
    put(ws, r, 1, "Deals with more than one sponsor are credited in full to each sponsor, so the sponsor totals exceed the deal count and commitments.", F(i=True, s=9))
    ns = wb["Notes"]
    n = len(deals)
    fixed = [
        f"Scope: {n} deals with status Closed in Raven.",
        "Date: the year and month in the Raven deal name. The first of the month is stored and shown as month and year, so order within a month is approximate. A deal's name may not match its IC memo date.",
        "Amount: loan amount as stored in Raven, assumed to be in $ millions. Blank means Raven shows 0.",
        "Sponsor: the sponsor label as Raven shows it. Columns G to J split it into separate sponsors, with abbreviations expanded to the firm name.",
        "Sponsor summary: every sponsor named on a deal is credited in full with that deal and its commitment. Deals with more than one sponsor therefore appear under each of them, and sponsor totals exceed the deal count and total commitments. The memo row shows the true totals.",
        "Commitments in the summary tables are the Amount column, so deals with blank amounts count as deals but add nothing to commitments.",
        "GICS Sector: one of the 11 GICS sectors, assigned from the company's principal business as described in the deal summary. Business is a short description from the same source.",
    ] + cfg.get("notes", [])
    for i, t in enumerate(fixed):
        c = ns.cell(3 + i, 1, t); c.font = F(); c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(out)
    print(f"wrote {out}: {n} deals, {len(sponsors)} sponsors")

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
