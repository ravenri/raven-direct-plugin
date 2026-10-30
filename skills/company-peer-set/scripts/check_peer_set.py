#!/usr/bin/env python3
"""Pre-delivery check for a Peer Set Word document. Run after building, before delivering.

Usage: python check_peer_set.py Peer_Set.docx
Exit code 1 if any hard failure. Prints every heading so they can be proofread as text.
"""
import re
import sys
from docx import Document

HARD = [
    (r"—", "em dash"),
    (r"\[\d+(?:,\s*\d+)*\]", "raw source bracket"),
    (r"\[\[n\d+\]\]", "number tag"),
    (r"\bc\d{1,3}\b", "internal source ID"),
    (r"\bQ\d{3,4}\b", "prompt ID"),
    (r"(?i)\bthe vault\b|\bvault\b", "vault reference"),
    (r"(?i)\bthe tile\b|\bcomp_analysis\b|\bdiligence_summary\b|\bthe leaf\b|\bthe family\b", "system jargon"),
    (r"(?i)\bAI summary\b|\bthe reads?\b", "system reference"),
    (r"(?i)\b(best-in-class|robust|compelling|well-positioned|world-class|market-leading|premier|stellar|poised to)\b", "banned hype word"),
    (r"(?i)\b(favour|colour|organis|analys(e|ed|ing)\b|behaviour|licence|centre)", "British spelling"),
    (r"(?i)\barguably\b|\bquite\b|\bvery\b|\btension\b|\bproven track record\b|\btranslates to\b|\bcorroborated\b", "banned word or phrase"),
    (r"(?i)\bnot stated\b", "'not stated' (check the stored analysis before using)"),
]
BAD_HEADINGS = {"overview", "assessment", "the read", "the take", "summary", "synthesis", "conclusion", "analysis"}
SOFT = [
    (r"(?i)\bsignificant(ly)?\b", "'significant' without a number"),
    (r"(?i)\bstrong(er|ly)?\b", "'strong' without a benchmark"),
    (r"(?i)\bkey\b", "'key' as filler"),
]


def all_text(doc):
    out = []
    for p in doc.paragraphs:
        out.append((p.style.name, p.text))
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                out.append(("table", c.text))
    for s in doc.sections:
        for p in s.footer.paragraphs:
            out.append(("footer", p.text))
    return out


def main(path):
    doc = Document(path)
    items = all_text(doc)
    fails, warns = [], []
    for style, text in items:
        for pat, label in HARD:
            for m in re.finditer(pat, text):
                fails.append(f"{label}: ...{text[max(0, m.start()-30):m.end()+30]}...")
        for pat, label in SOFT:
            for m in re.finditer(pat, text):
                warns.append(f"{label}: ...{text[max(0, m.start()-30):m.end()+30]}...")
    heads = [(s, t) for s, t in items if s.startswith("Heading") or s == "Title"]
    for s, t in heads:
        if t.strip().lower() in BAD_HEADINGS:
            fails.append(f"meta-label heading: {t}")
        if "?" in t or ":" in t:
            fails.append(f"heading is a question or has a colon: {t}")
    text_all = "\n".join(t for _, t in items)
    # structural checks
    n_signals = text_all.count("Deal says.")
    n_filings = text_all.count("Filings reveal.")
    if n_signals != n_filings:
        fails.append(f"Filing Signals mismatch: {n_signals} 'Deal says' vs {n_filings} 'Filings reveal'")
    n_items = sum(1 for s, t in items if s == "Heading 2")
    n_sources = sum(1 for s, t in items if t.startswith("Source:"))
    if n_sources < n_items:
        fails.append(f"{n_items} items but only {n_sources} Source lines")
    if [t for _, t in heads][1:2] != ["Executive Summary"]:
        fails.append("Executive Summary must be the first section after the title")
    for need in ("Executive Summary", "Peer Benchmarking"):
        if need not in [t for _, t in heads]:
            fails.append(f"missing section: {need}")
    print("HEADINGS (proofread these as text):")
    for s, t in heads:
        print(f"  [{s}] {t}")
    print(f"\nItems: {n_items}  Source lines: {n_sources}  Tables: {len(doc.tables)}")
    if warns:
        print("\nREVIEW (soft flags):")
        for w in warns:
            print("  -", w)
    if fails:
        print("\nFAIL:")
        for f in fails:
            print("  -", f)
        sys.exit(1)
    print("\nPASS: no hard failures")


if __name__ == "__main__":
    main(sys.argv[1])
