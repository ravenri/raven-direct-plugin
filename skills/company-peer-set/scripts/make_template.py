"""Regenerates assets/peer_set_template.docx (layout and style guideline for the Peer Set document).

Run: python make_template.py ../assets/peer_set_template.docx
The template holds the page setup, fonts, colors and named styles that build_peer_set.py uses.
Its body shows each element once with placeholder wording only. It is a guideline for layout
and style, not the canonical form of any output. To change the look, edit the styles in the
.docx in Word (or here) and keep the style names.
"""
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Arial"
INK = RGBColor(0x24, 0x2D, 0x27)      # dark green, headings
GREY = RGBColor(0x5A, 0x63, 0x6A)     # secondary text


def set_font(style, size, bold=False, italic=False, color=None):
    style.font.name = FONT
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color


def add_style(doc, name, size, bold=False, italic=False, color=None, before=0, after=3, keep_next=False, base="Normal"):
    st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = doc.styles[base]
    set_font(st, size, bold, italic, color)
    pf = st.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.keep_with_next = keep_next
    pf.line_spacing = 1.08
    return st


def bottom_border(style, color="242D27", sz=6):
    pPr = style.element.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(sz)); b.set(qn("w:space"), "2"); b.set(qn("w:color"), color)
    pb.append(b)
    pPr.append(pb)


def main(out):
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(11), Inches(8.5)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(0.75))
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)

    n = doc.styles["Normal"]
    set_font(n, 9.5)
    n.paragraph_format.space_after = Pt(4)

    # Built-in heading styles carry the outline levels
    h1 = doc.styles["Heading 1"]; set_font(h1, 13, True, False, INK)
    h1.paragraph_format.space_before = Pt(14); h1.paragraph_format.space_after = Pt(6); h1.paragraph_format.keep_with_next = True
    bottom_border(h1)
    h2 = doc.styles["Heading 2"]; set_font(h2, 10.5, True, False, INK)
    h2.paragraph_format.space_before = Pt(9); h2.paragraph_format.space_after = Pt(1); h2.paragraph_format.keep_with_next = True
    h3 = doc.styles["Heading 3"]; set_font(h3, 10, True, False, INK)
    h3.paragraph_format.space_before = Pt(10); h3.paragraph_format.space_after = Pt(3); h3.paragraph_format.keep_with_next = True
    t = doc.styles["Title"]; set_font(t, 22, True, False, INK)
    t.paragraph_format.space_after = Pt(2)
    pPr = t.element.get_or_add_pPr()
    for el in pPr.findall(qn("w:pBdr")):
        pPr.remove(el)

    add_style(doc, "PS Subtitle", 10.5, color=GREY, after=2)
    add_style(doc, "PS Meta", 8.5, color=GREY, after=2, keep_next=True)
    add_style(doc, "PS Body", 9.5, after=4)
    add_style(doc, "PS Label", 9.5, after=4)
    add_style(doc, "PS Note", 8, color=GREY, after=2)
    add_style(doc, "PS Source", 8, italic=True, color=GREY, after=6)
    add_style(doc, "PS Table", 8, after=0)
    add_style(doc, "PS Bullet", 9.5, after=3)

    # placeholder body that shows each element once
    doc.add_paragraph("Peer Set", style="Title")
    doc.add_paragraph("Company name | Deal name", style="PS Subtitle")
    doc.add_paragraph("Confirmed public peers: names | Generated date", style="PS Meta")
    doc.add_paragraph("Peer Benchmarking", style="Heading 1")
    doc.add_paragraph("Section heading names the content, one plain noun phrase.", style="PS Body")
    doc.add_paragraph("Item label", style="Heading 2")
    doc.add_paragraph("Tag or severity line", style="PS Meta")
    p = doc.add_paragraph(style="PS Label")
    r = p.add_run("Credit implication. "); r.bold = True
    p.add_run("Placeholder wording for the implication line.")
    doc.add_paragraph("Source: Issuer filing, fiscal period.", style="PS Source")
    doc.save(out)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "peer_set_template.docx")
