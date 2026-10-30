#!/usr/bin/env python3
"""Builds the Peer Set Word document from a build-input file, using assets/peer_set_template.docx.

Usage:
  python build_peer_set.py --input build_input.json --template ../assets/peer_set_template.docx --out Peer_Set.docx

The builder only lays out text it is given. It never computes a figure. It enforces the length
budgets below before building (use --force to build anyway) and applies **bold** markers in text.
Section order: Executive Summary, then, as on the platform page, Peer Benchmarking, Industry Developments, Peer Insights,
Filing Signals, Capital Structure. Empty sections are skipped.

Input shape (all text plain; **double asterisks** mark bold figures):
  title, company, deal, meta_line
  executive_summary: {lead, points[3 to 5]}  (written last, from the finished sections; 260 words total)
  benchmarking: caveats[], columns[], rows[{cells[], kind: target|peer|median}], row_notes[], data_notes[], takeaways[]
  industry_developments[]: {label, tag, body, credit_implication, sources}
  peer_insights[]: {label, body, credit_implication, deal_note, sources}
  filing_signals[]: {label, severity, classification, deal_says, filings_reveal, sources}
  capital_structure: {covenant_narrative, structural_caveat, target_reference, leverage_basis,
                      blocks[{company, as_of, rows[[facility, outstanding, rate, maturity, security, key_terms]], source}]}
"""
import argparse
import json
import re
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Length budgets, in words. Summaries are capped; Filing Signals are kept near full length.
BUDGET = {
    "development_body": 80,
    "insight_body": 80,
    "credit_implication": 55,
    "deal_note": 55,
    "takeaway": 80,
    "signal_text": 130,
    "narrative": 120,
    "cell": 40,
    "caveat": 60,
    "exec_total": 260,
    "exec_point": 55,
}
BENCH_WIDTHS = [1.7, 0.8, 0.8, 0.8, 0.85, 0.7, 0.85, 1.3, 0.85, 0.85]   # inches, sums to 9.5
CAP_COLUMNS = ["Facility", "Outstanding", "Rate", "Maturity", "Security", "Key terms"]
CAP_WIDTHS = [1.7, 1.8, 1.1, 1.3, 1.0, 2.6]                             # inches, sums to 9.5


def words(text):
    return len(re.findall(r"\S+", text or ""))


def validate(d):
    errs = []

    def cap(name, text, limit):
        n = words(text)
        if n > limit:
            errs.append(f"{name}: {n} words, limit {limit}")

    ex = d.get("executive_summary") or {}
    if not ex:
        errs.append("executive_summary: missing (required, written last from the finished sections)")
    else:
        pts = ex.get("points", [])
        if not 3 <= len(pts) <= 5:
            errs.append(f"executive_summary: {len(pts)} points, needs 3 to 5")
        total = words(ex.get("lead", "")) + sum(words(p) for p in pts)
        if total > BUDGET["exec_total"]:
            errs.append(f"executive_summary: {total} words, limit {BUDGET['exec_total']}")
        for i, p in enumerate(pts):
            cap(f"executive_summary point {i+1}", p, BUDGET["exec_point"])
    b = d.get("benchmarking") or {}
    for i, t in enumerate(b.get("takeaways", [])):
        cap(f"takeaway {i+1}", t, BUDGET["takeaway"])
    for i, t in enumerate(b.get("caveats", [])):
        cap(f"caveat {i+1}", t, BUDGET["caveat"])
    for it in d.get("industry_developments", []):
        cap(f"development '{it['label']}' body", it["body"], BUDGET["development_body"])
        cap(f"development '{it['label']}' credit implication", it["credit_implication"], BUDGET["credit_implication"])
        if not it.get("sources"):
            errs.append(f"development '{it['label']}': missing sources")
    for it in d.get("peer_insights", []):
        cap(f"insight '{it['label']}' body", it["body"], BUDGET["insight_body"])
        cap(f"insight '{it['label']}' credit implication", it["credit_implication"], BUDGET["credit_implication"])
        cap(f"insight '{it['label']}' deal note", it["deal_note"], BUDGET["deal_note"])
        if not it.get("deal_note"):
            errs.append(f"insight '{it['label']}': missing deal note")
        if not it.get("sources"):
            errs.append(f"insight '{it['label']}': missing sources")
    for it in d.get("filing_signals", []):
        if not it.get("deal_says"):
            errs.append(f"signal '{it['label']}': missing Deal Says (never drop the deal-silent note)")
        cap(f"signal '{it['label']}' filings reveal", it["filings_reveal"], BUDGET["signal_text"])
        cap(f"signal '{it['label']}' deal says", it.get("deal_says", ""), BUDGET["signal_text"])
        if not it.get("sources"):
            errs.append(f"signal '{it['label']}': missing sources")
    cs = d.get("capital_structure") or {}
    for k in ("covenant_narrative", "structural_caveat", "target_reference", "leverage_basis"):
        cap(f"capital structure {k}", cs.get(k, ""), BUDGET["narrative"])
    for blk in cs.get("blocks", []):
        if not blk.get("source"):
            errs.append(f"capital structure block '{blk['company']}': missing source")
        for row in blk["rows"]:
            if len(row) != len(CAP_COLUMNS):
                errs.append(f"capital structure '{blk['company']}': row has {len(row)} cells, needs {len(CAP_COLUMNS)}")
            for c in row:
                cap(f"cell in '{blk['company']}'", c, BUDGET["cell"])
    return errs


# ---------- low-level helpers ----------
def runs(par, text, size=None, color=None, italic=None):
    """Adds text to a paragraph, honoring **bold** markers."""
    for i, chunk in enumerate(re.split(r"\*\*", text)):
        if not chunk:
            continue
        r = par.add_run(chunk)
        if i % 2 == 1:
            r.bold = True
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
        if italic is not None:
            r.italic = italic
    return par


def para(doc, text, style):
    p = doc.add_paragraph(style=style)
    runs(p, text)
    return p


def labeled(doc, label, text, style="PS Label"):
    p = doc.add_paragraph(style=style)
    r = p.add_run(label + " ")
    r.bold = True
    runs(p, text)
    p.paragraph_format.keep_together = True
    return p


def set_cell_borders(cell, top=None, bottom=None):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge, spec in (("top", top), ("bottom", bottom)):
        if spec:
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(spec[0])); el.set(qn("w:color"), spec[1]); el.set(qn("w:space"), "0")
            borders.append(el)
    tcPr.append(borders)


def cell_margins(table, top=40, bottom=40, left=60, right=60):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{k}"); el.set(qn("w:w"), str(v)); el.set(qn("w:type"), "dxa"); m.append(el)
    tblPr.append(m)


def make_table(doc, columns, rows, widths, kinds=None, keep_together=False):
    t = doc.add_table(rows=1, cols=len(columns))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    cell_margins(t)
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for gc, wd in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(wd * 1440)))
    GREY = "9CA3AF"; INK = "242D27"
    hdr = t.rows[0]
    trPr = hdr._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for i, c in enumerate(columns):
        cell = hdr.cells[i]
        cell.width = Inches(widths[i])
        p = cell.paragraphs[0]; p.style = doc.styles["PS Table"]
        p.paragraph_format.keep_with_next = True
        r = p.add_run(c); r.bold = True
        set_cell_borders(cell, bottom=(8, INK))
    for ri, row in enumerate(rows):
        kind = (kinds or [None] * len(rows))[ri]
        cells = t.add_row().cells
        trPr = t.rows[-1]._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
        for i, val in enumerate(row):
            cells[i].width = Inches(widths[i])
            p = cells[i].paragraphs[0]; p.style = doc.styles["PS Table"]
            if keep_together:
                p.paragraph_format.keep_with_next = True
            runs(p, val if val not in (None, "") else "–")
            if kind in ("target", "median"):
                for r in p.runs:
                    r.bold = True
            top = (6, INK) if kind == "median" else None
            set_cell_borders(cells[i], top=top, bottom=(4, GREY))
    doc.add_paragraph(style="PS Note")
    return t


def clear_body(doc):
    body = doc.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)


def footer(doc, left_text):
    sec = doc.sections[0]
    f = sec.footer
    f.is_linked_to_previous = False
    p = f.paragraphs[0] if f.paragraphs else f.add_paragraph()
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.style = doc.styles["PS Note"]
    p.add_run(left_text + "    Page ")
    r = p.add_run()
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)


# ---------- sections ----------
def sec_exec(doc, ex):
    if not ex:
        return
    doc.add_paragraph("Executive Summary", style="Heading 1")
    para(doc, ex["lead"], "PS Body")
    for p_ in ex["points"]:
        para(doc, p_, "PS Bullet").style = doc.styles["List Number"]


def sec_benchmarking(doc, b):
    if not b:
        return
    doc.add_paragraph("Peer Benchmarking", style="Heading 1")
    for c in b.get("caveats", []):
        para(doc, c, "PS Body")
    rows = [r["cells"] for r in b["rows"]]
    kinds = [r.get("kind") for r in b["rows"]]
    make_table(doc, b["columns"], rows, BENCH_WIDTHS[: len(b["columns"])], kinds)
    for n in b.get("row_notes", []):
        para(doc, n, "PS Note")
    if b.get("data_notes"):
        doc.add_paragraph("Data notes", style="Heading 3")
        for n in b["data_notes"]:
            para(doc, n, "PS Bullet").style = doc.styles["List Bullet"]
    if b.get("takeaways"):
        doc.add_paragraph("Takeaways", style="Heading 3")
        for t in b["takeaways"]:
            p = para(doc, t, "PS Bullet")
            p.style = doc.styles["List Bullet"]


def sec_developments(doc, items):
    if not items:
        return
    doc.add_paragraph("Industry Developments", style="Heading 1")
    for it in items:
        doc.add_paragraph(it["label"], style="Heading 2")
        doc.add_paragraph(it["tag"], style="PS Meta")
        para(doc, it["body"], "PS Body")
        labeled(doc, "Credit implication.", it["credit_implication"])
        para(doc, "Source: " + it["sources"], "PS Source")


def sec_insights(doc, items):
    if not items:
        return
    doc.add_paragraph("Peer Insights", style="Heading 1")
    for it in items:
        doc.add_paragraph(it["label"], style="Heading 2")
        para(doc, it["body"], "PS Body")
        labeled(doc, "Credit implication.", it["credit_implication"])
        labeled(doc, "Deal:", it["deal_note"])
        para(doc, "Source: " + it["sources"], "PS Source")


def sec_signals(doc, items):
    if not items:
        return
    doc.add_paragraph("Filing Signals", style="Heading 1")
    for it in items:
        doc.add_paragraph(it["label"], style="Heading 2")
        doc.add_paragraph(f"Severity: {it['severity']}  |  Classification: {it['classification']}", style="PS Meta")
        labeled(doc, "Deal says.", it["deal_says"])
        labeled(doc, "Filings reveal.", it["filings_reveal"])
        para(doc, "Source: " + it["sources"], "PS Source")


def sec_capital(doc, cs):
    if not cs:
        return
    doc.add_paragraph("Capital Structure", style="Heading 1")
    for label, key in (("Covenants.", "covenant_narrative"), ("Comparability.", "structural_caveat"),
                       ("Target reference.", "target_reference"), ("Leverage basis.", "leverage_basis")):
        if cs.get(key):
            labeled(doc, label, cs[key])
    for blk in cs.get("blocks", []):
        doc.add_paragraph(f"{blk['company']}, as of {blk['as_of']}" if blk.get("as_of") else blk["company"], style="Heading 3")
        make_table(doc, CAP_COLUMNS, blk["rows"], CAP_WIDTHS, keep_together=len(blk["rows"]) <= 12)
        para(doc, "Source: " + blk["source"], "PS Source")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--template", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--force", action="store_true", help="build even if a length budget is exceeded")
    a = ap.parse_args()
    d = json.load(open(a.input, encoding="utf-8"))
    errs = validate(d)
    if errs:
        print("Length and completeness check:")
        for e in errs:
            print("  -", e)
        if not a.force:
            sys.exit(1)
    doc = Document(a.template)
    for need in ("PS Body", "PS Label", "PS Meta", "PS Note", "PS Source", "PS Table", "PS Bullet", "PS Subtitle"):
        if need not in [s.name for s in doc.styles]:
            sys.exit(f"Template is missing style '{need}'")
    clear_body(doc)
    doc.add_paragraph(d.get("title", "Peer Set"), style="Title")
    doc.add_paragraph(f"{d['company']} | {d['deal']}", style="PS Subtitle")
    doc.add_paragraph(d["meta_line"], style="PS Meta")
    sec_exec(doc, d.get("executive_summary"))
    sec_benchmarking(doc, d.get("benchmarking"))
    sec_developments(doc, d.get("industry_developments"))
    sec_insights(doc, d.get("peer_insights"))
    sec_signals(doc, d.get("filing_signals"))
    sec_capital(doc, d.get("capital_structure"))
    footer(doc, f"Peer Set | {d['company']}")
    doc.core_properties.title = f"Peer Set, {d['company']}"
    doc.save(a.out)
    print("Built", a.out)


if __name__ == "__main__":
    main()
